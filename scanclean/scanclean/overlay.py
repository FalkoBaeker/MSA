"""Maximale Treue: Hintergrundbild + editierbare Textebene.

Idee:
  * Der Hintergrund ist eine weiße Seite, auf der nur die *bewahrten Grafiken*
    liegen – Logo/Briefkopf, Tabellen (als Bild) und Trennlinien. Der gesamte
    Originaltext ist entfernt; auch das CamScanner-Wasserzeichen verschwindet.
  * Darüber liegt der erkannte Text als *absolut positionierte Textfelder* an
    den Originalkoordinaten. Dieser Text ist voll editierbar.

Ergebnis: sieht praktisch aus wie das Original, lässt sich aber bearbeiten.

Die Textfelder werden als OOXML-DrawingML (verankerte Shapes) erzeugt, weil
python-docx dafür keine eigene API hat. Das DOCX bleibt voll Word-kompatibel.
"""

from __future__ import annotations

from xml.sax.saxutils import escape

import cv2
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Emu
from docx.oxml import parse_xml

from .layout import PageModel, Line, px_to_pt

EMU_PER_INCH = 914400
DEFAULT_FONT = "Arial"


def _emu(px: float, dpi: int) -> int:
    return int(round(px / dpi * EMU_PER_INCH))


def make_background(model: PageModel, cleaned: np.ndarray) -> np.ndarray:
    """Weiße Seite mit nur den bewahrten Grafiken (Logo, Tabellen, Linien)."""
    bg = np.full_like(cleaned, 255)
    for g in model.graphics + model.tables:
        bg[g.top:g.bottom, g.left:g.right] = cleaned[g.top:g.bottom, g.left:g.right]
    for h in model.hlines:
        cv2.line(bg, (h.x1, h.y1), (h.x2, h.y2), 0, 2)
    return bg


def _background_xml(rid: str, cx: int, cy: int, shape_id: int) -> str:
    return (
        '<w:drawing '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<wp:anchor behindDoc="1" distT="0" distB="0" distL="0" distR="0" '
        'simplePos="0" locked="0" layoutInCell="1" allowOverlap="1" relativeHeight="0">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{shape_id}" name="bg{shape_id}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{shape_id}" name="bg{shape_id}"/>'
        '<pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing>'
    )


def _textbox_xml(line: Line, dpi: int, sz_halfpt: int, shape_id: int,
                 top_px: float) -> str:
    x = _emu(line.left, dpi)
    y = _emu(top_px, dpi)
    # Box etwas breiter/höher als gemessen, Text läuft ohne Umbruch (wrap none)
    w = _emu((line.right - line.left) + dpi // 2, dpi)
    h = _emu((line.bottom - line.top) + dpi // 5, dpi)
    bold = "<w:b/>" if line.bold else ""
    text = escape(line.text)
    align = ""
    if line.align == "right":
        align = '<w:jc w:val="right"/>'
    elif line.align == "center":
        align = '<w:jc w:val="center"/>'
    return (
        '<w:drawing '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<wp:anchor behindDoc="0" distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" locked="0" layoutInCell="1" allowOverlap="1" relativeHeight="{1000 + shape_id}">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{w}" cy="{h}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{shape_id}" name="t{shape_id}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData '
        'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/>'
        f'<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
        '<wps:txbx><w:txbxContent><w:p><w:pPr>'
        '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        f'{align}</w:pPr>'
        f'<w:r><w:rPr><w:rFonts w:ascii="{DEFAULT_FONT}" w:hAnsi="{DEFAULT_FONT}"/>'
        f'<w:sz w:val="{sz_halfpt}"/><w:szCs w:val="{sz_halfpt}"/>{bold}</w:rPr>'
        f'<w:t xml:space="preserve">{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        '<wps:bodyPr rot="0" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" '
        'anchor="t" anchorCtr="0"/></wps:wsp></a:graphicData></a:graphic>'
        '</wp:anchor></w:drawing>'
    )


def _font_halfpts(lines: list[Line], dpi: int, body_pt: float) -> list[int]:
    """Schriftgröße (in Halbpunkten) je Zeile – einheitlich wie im Original.

    Das Original verwendet praktisch *eine* Textgröße; Überschriften sind nur
    fett, nicht größer. Wir setzen daher fast alle Zeilen auf die dominante
    Größe und vergrößern nur deutlich größere Titel moderat. Überlappungen
    werden später durch leichtes Verschieben (Entzerrung) gelöst, nicht durch
    Schrumpfen – so bleibt das Schriftbild ruhig.
    """
    # Tesseract misst die Wort-Höhe enger als die Schriftgröße -> hochskalieren.
    # Eine *einheitliche* Größe (wie im Original) ist robuster als pro-Zeile
    # gemessene Größen, die bei Ober-/Unterlängen leicht ausreißen.
    body_eff = body_pt * 1.2
    body_px = body_eff * dpi / 72.0
    big_px = body_px * 1.7                       # nur klar größere Titel
    sizes: list[int] = []
    for ln in lines:
        if ln.raw_h > big_px:
            pt = min(px_to_pt(ln.raw_h, dpi) * 1.15, body_eff * 1.6)
        else:
            pt = body_eff
        sizes.append(max(8, int(round(pt * 2))))
    return sizes


def _decollide_tops(lines: list[Line], sizes: list[int], dpi: int) -> list[float]:
    """Schiebt sich überlappende Textzeilen minimal nach unten.

    Der Hintergrund ist unter dem Text weiß, daher darf der Text leicht
    verschoben werden, ohne die Grafiken (Logo/Tabelle/Linien) zu treffen.
    Es wird nur gegen die vorige Zeile geschoben und nur, wenn sie sich
    waagerecht überlappen (damit z. B. rechtsbündige Elemente unberührt bleiben).
    """
    order = sorted(range(len(lines)), key=lambda i: (lines[i].top, lines[i].left))
    tops = [float(ln.top) for ln in lines]
    prev_idx = None
    for idx in order:
        ln = lines[idx]
        if prev_idx is not None:
            prev = lines[prev_idx]
            h_overlap = (ln.left < prev.right and prev.left < ln.right)
            prev_h = sizes[prev_idx] / 2.0 * dpi / 72.0 * 1.18
            min_top = tops[prev_idx] + prev_h
            if h_overlap and tops[idx] < min_top:
                tops[idx] = min_top
        prev_idx = idx
    return tops


def _add_page(doc, section, model: PageModel, cleaned: np.ndarray,
              shape_base: int) -> int:
    section.page_width = Emu(_emu(model.width, model.dpi))
    section.page_height = Emu(_emu(model.height, model.dpi))
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin",
                 "header_distance", "footer_distance"):
        setattr(section, attr, Emu(0))

    bg = make_background(model, cleaned)
    ok, buf = cv2.imencode(".png", bg)
    import io
    rid, _ = doc.part.get_or_add_image(io.BytesIO(buf.tobytes()))
    cx = _emu(model.width, model.dpi)
    cy = _emu(model.height, model.dpi)

    para = doc.add_paragraph()
    run = para.add_run()
    run._r.append(parse_xml(_background_xml(rid, cx, cy, shape_base)))

    sizes = _font_halfpts(model.lines, model.dpi, model.body_pt)
    tops = _decollide_tops(model.lines, sizes, model.dpi)
    for i, line in enumerate(model.lines):
        r = para.add_run()
        r._r.append(parse_xml(_textbox_xml(line, model.dpi, sizes[i],
                                           shape_base + 1 + i, tops[i])))
    return shape_base + 1 + len(model.lines)


def build_overlay_document(pages: list[tuple[PageModel, np.ndarray]],
                           out_path: str) -> None:
    """Baut das DOCX (Hintergrundbild + Textebene) über alle Seiten."""
    doc = Document()
    shape_base = 1
    for i, (model, cleaned) in enumerate(pages):
        if i == 0:
            section = doc.sections[0]
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
        shape_base = _add_page(doc, section, model, cleaned, shape_base) + 10
    doc.save(out_path)
