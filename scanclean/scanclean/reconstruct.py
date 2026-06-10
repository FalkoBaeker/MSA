"""Rekonstruktion: aus dem Seitenmodell ein editierbares DOCX bauen.

Ziel ist ein frisches, weißes Dokument mit echtem Text an möglichst
originalgetreuer Position, echten Trennlinien und dem Logo als Bild.
Der Text bleibt voll bearbeitbar (Tippfehler, Unterstreichungen ...).
"""

from __future__ import annotations

import io

import cv2
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

from .layout import PageModel, Line, px_to_pt

# Standardfont: serifenlos, nah an der valera-Schrift (Arial/Helvetica-Klasse)
DEFAULT_FONT = "Arial"


def _set_paragraph_border_bottom(paragraph) -> None:
    """Zeichnet eine echte Linie unter einen (leeren) Absatz."""
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    pPr.append(borders)


def _add_line_paragraph(doc, line: Line, dpi: int, indent_step_px: int,
                        space_before_pt: float) -> None:
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(max(0.0, space_before_pt))
    pf.line_spacing = 1.05

    if line.align == "right":
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif line.align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if line.indent_level > 0:
            pf.left_indent = Pt(px_to_pt(line.indent_level * indent_step_px, dpi))

    run = para.add_run(line.text)
    run.font.name = DEFAULT_FONT
    run.font.size = Pt(line.font_pt)
    run.font.bold = line.bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_logo(doc, cleaned: np.ndarray, graphic, dpi: int) -> None:
    crop = cleaned[graphic.top:graphic.bottom, graphic.left:graphic.right]
    if crop.size == 0:
        return
    ok, buf = cv2.imencode(".png", crop)
    if not ok:
        return
    stream = io.BytesIO(buf.tobytes())
    width_in = (graphic.right - graphic.left) / dpi
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run()
    run.add_picture(stream, width=Inches(width_in))


def _add_region_image(doc, cleaned: np.ndarray, region, dpi: int,
                      space_before_pt: float) -> None:
    """Fügt einen Bildausschnitt (z. B. Tabelle) an Originalbreite ein."""
    crop = cleaned[region.top:region.bottom, region.left:region.right]
    if crop.size == 0:
        return
    ok, buf = cv2.imencode(".png", crop)
    if not ok:
        return
    stream = io.BytesIO(buf.tobytes())
    width_in = min(6.8, (region.right - region.left) / dpi)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(max(0.0, space_before_pt))
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run()
    run.add_picture(stream, width=Inches(width_in))


def _configure_section(section, model: PageModel) -> int:
    """Setzt A4 und die Ränder einer Seite; liefert den oberen Rand in px."""
    section.page_width = Inches(8.27)    # A4
    section.page_height = Inches(11.69)
    # linker Rand robust aus den Zeilen (10%-Perzentil der Startpositionen)
    lefts = sorted(ln.left for ln in model.lines)
    margin_left_px = lefts[len(lefts) // 10] if lefts else 0
    # oberster Rand inkl. Logo, damit der Briefkopf ganz oben sitzt
    tops = [ln.top for ln in model.lines] + [g.top for g in model.graphics]
    margin_top_px = min(tops) if tops else 0
    section.left_margin = Inches(max(0.4, margin_left_px / model.dpi))
    section.right_margin = Inches(0.4)
    section.top_margin = Inches(max(0.4, margin_top_px / model.dpi))
    section.bottom_margin = Inches(0.4)
    return margin_top_px


def _render_page(doc, model: PageModel, cleaned: np.ndarray,
                 margin_top_px: int) -> None:
    """Schreibt den Inhalt einer Seite (Logo, Text, Tabellen, Linien)."""
    # Logo(s) zuerst (liegen oben)
    for g in model.graphics:
        _add_logo(doc, cleaned, g, model.dpi)

    hlines = sorted(model.hlines, key=lambda h: h.y1)
    hl_idx = 0
    prev_bottom = margin_top_px
    # eine "Zeilenhöhe" in Pixel, um Leerzeilen-Abstände zu erkennen
    body_line_px = model.body_pt * model.dpi / 72.0

    # Zeilen und Tabellen-Bilder in Lesereihenfolge (nach y) mischen
    elements: list[tuple[int, str, object]] = []
    for line in model.lines:
        elements.append((line.top, "line", line))
    for table in model.tables:
        elements.append((table.top, "table", table))
    elements.sort(key=lambda e: e[0])

    for top, kind, obj in elements:
        # Trennlinien einfügen, die vor diesem Element liegen
        while hl_idx < len(hlines) and hlines[hl_idx].y1 < top:
            sep = doc.add_paragraph()
            sep.paragraph_format.space_before = Pt(2)
            sep.paragraph_format.space_after = Pt(2)
            _set_paragraph_border_bottom(sep)
            hl_idx += 1

        gap_px = top - prev_bottom
        blanks = max(0, min(3, round(gap_px / body_line_px) - 1)) if body_line_px else 0
        space_before = blanks * model.body_pt

        if kind == "table":
            _add_region_image(doc, cleaned, obj, model.dpi, space_before)
            prev_bottom = obj.bottom
        else:
            _add_line_paragraph(doc, obj, model.dpi, model.indent_step_px,
                                space_before)
            prev_bottom = obj.bottom


def build_document(pages: list[tuple[PageModel, np.ndarray]],
                   out_path: str) -> None:
    """Baut aus mehreren Seitenmodellen *ein* DOCX (eine Sektion pro Seite)."""
    doc = Document()
    for i, (model, cleaned) in enumerate(pages):
        if i > 0:
            doc.add_section(WD_SECTION.NEW_PAGE)
        margin_top_px = _configure_section(doc.sections[-1], model)
        _render_page(doc, model, cleaned, margin_top_px)
    doc.save(out_path)


def build_docx(model: PageModel, cleaned: np.ndarray, out_path: str) -> None:
    """Einzelseite – Bequemlichkeits-Wrapper um build_document()."""
    build_document([(model, cleaned)], out_path)
